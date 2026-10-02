"""Failures in plain words, the off switches, error pages and one connection per page."""

import io
import logging
import time
import zipfile

import pytest
from openpyxl import load_workbook

from leadgen import calls, config, daily, marks, pipeline, saved, stats, store, web
from leadgen.export import to_xlsx_bytes
from leadgen.http import HttpError
from leadgen.models import Lead
from leadgen.pipeline import PipelineError, RunResult, SearchParams
from leadgen.progress import MAP, Step
from leadgen.scoring import score_lead
from leadgen.sources import osm

TECHNICAL = ("http://", "https://", "Error", "Exception", "<html", "HTTP ")


def _wait(client, job):
    for _ in range(100):
        body = client.get(f"/status/{job}").get_json()
        if body["state"] != "running":
            return body
        time.sleep(0.05)
    raise AssertionError("the search never finished")


def _lead(name="Costco", sid="c", **kw):
    kw.setdefault("lat", 40.72)
    kw.setdefault("lon", -111.9)
    kw.setdefault("raw_categories", ["shop=wholesale"])
    lead = Lead(name=name, source="osm", source_id=sid, **kw)
    score_lead(lead, config.DEFAULT_KEYWORDS)
    return lead


# ---- a failed search: fast, plain, logged, and in the history

def test_map_data_step_has_one_overall_deadline(monkeypatch):
    clock = {"now": 1000.0}
    monkeypatch.setattr(osm.time, "monotonic", lambda: clock["now"])
    timeouts = []

    def slow_failure(method, url, **kw):
        timeouts.append(kw["timeout"])
        assert kw["retries"] == 1
        clock["now"] += kw["timeout"][1]          # the server hangs as long as allowed
        raise HttpError(f"{url}: ConnectionResetError(104, 'reset')")
    monkeypatch.setattr(osm, "request_json", slow_failure)
    with pytest.raises(osm.SourceError):
        osm.search(40.76, -111.89, 5)
    # The whole query, then its quarters, until the first round's time is used up, then
    # the catch-up rounds, each with its own time; no part waits longer than a part's
    # share (two parts run at once, so one may overrun each round).
    rounds = config.OVERPASS_RETRY_ROUNDS
    assert clock["now"] - 1000 <= (config.OVERPASS_DEADLINE_SECONDS
                                   + rounds * config.OVERPASS_RETRY_SECONDS
                                   + (rounds + 1) * config.OVERPASS_PART_SECONDS)
    assert all(t[1] <= config.OVERPASS_PART_SECONDS for t in timeouts)
    assert len(timeouts) <= 5 * (rounds + 1)
    assert f"[timeout:{config.OVERPASS_PART_SECONDS}]" in osm.build_query(40.76, -111.89, 30)


def test_quick_mirror_failures_try_the_next_one(monkeypatch):
    tried = []

    def fail_then_answer(method, url, **kw):
        tried.append(url)
        if len(tried) < 3:
            raise HttpError(f"{url} returned HTTP 504: <html><title>504 Gateway Time-out")
        return {"elements": []}
    monkeypatch.setattr(osm, "request_json", fail_then_answer)
    assert osm.search(40.76, -111.89, 5) == ([], [])
    assert tried == config.OVERPASS_ENDPOINTS[:3]


def test_a_hanging_mirror_does_not_stop_the_others(monkeypatch):
    """Mirror 1 fails at once, mirror 2 hangs longer than the whole deadline, mirror 3
    answers: the search succeeds with mirror 3's businesses, within the deadline."""
    import threading
    monkeypatch.setattr(config, "OVERPASS_DEADLINE_SECONDS", 8)
    monkeypatch.setattr(config, "OVERPASS_STAGGER_SECONDS", 0.5)
    released = threading.Event()
    tried = []

    def mirrors(method, url, **kw):
        tried.append(url)
        n = config.OVERPASS_ENDPOINTS.index(url)
        if n == 0:
            raise HttpError(f"{url}: ConnectionResetError(104, 'reset')")
        if n == 1:
            released.wait(30)                    # hangs past the deadline
            raise HttpError(f"{url}: ReadTimeout")
        return {"elements": [{"type": "node", "id": 7, "lat": 40.7, "lon": -111.9,
                              "tags": {"name": "Harmons", "shop": "supermarket"}}]}
    monkeypatch.setattr(osm, "request_json", mirrors)
    started = time.monotonic()
    try:
        leads, warnings = osm.search(40.76, -111.89, 5)
    finally:
        released.set()
    assert [l.name for l in leads] == ["Harmons"] and warnings == []
    assert time.monotonic() - started < config.OVERPASS_DEADLINE_SECONDS
    assert tried == config.OVERPASS_ENDPOINTS[:3]


def test_a_failed_search_is_plain_logged_and_kept_in_the_history(monkeypatch, caplog):
    def unreachable(*args, **kw):
        raise osm.SourceError("All OpenStreetMap (Overpass) servers failed: "
                              "HTTPSConnectionPool(host='overpass.kumi.systems') ProxyError(...)")
    monkeypatch.setattr(pipeline.osm, "search", unreachable)
    monkeypatch.setattr(pipeline, "geocode", lambda location, key: (40.76, -111.89, "Salt Lake City"))
    client = web.create_app().test_client()
    with caplog.at_level(logging.WARNING):
        job = client.post("/search", data={"location": "84101"}).get_json()["job_id"]
        body = _wait(client, job)
    assert body["state"] == "error"
    message = body["message"]
    assert "Couldn't reach the map data service" in message
    assert "Today's search was not used up" in message
    assert "You can try again now; if it fails again, try later today." in message
    assert not any(t in message for t in TECHNICAL)
    assert "ProxyError" in caplog.text                 # the detail is in the server log
    history = client.get("/searches").get_json()
    assert not history["used_today"]                    # the day is given back
    failed = history["searches"][0]
    assert failed["failed"] and "map data service" in failed["reason"]
    again = client.post("/search", data={"location": "84101"})
    assert again.status_code == 200
    # Wait for it, so the search doesn't carry on after this test's stand-ins are gone.
    assert _wait(client, again.get_json()["job_id"])["state"] == "error"


def test_partial_failures_read_as_plain_notes():
    assert not any(t in web.plain_warning(
        "Google 'grocery': the search failed (https://places.googleapis.com returned HTTP 500: "
        "<html>oops</html>); kept the results already found") for t in TECHNICAL)
    capped = web.plain_warning("Stopped at the 10-request Yelp cap: 3 Yelp searches not run, 0 "
                               "had more result pages. Raise the request cap (--max-requests).")
    assert "Yelp" in capped and "cap" not in capped and "--" not in capped
    assert web.plain_warning("a note") == "a note"


def test_pipeline_names_the_source_it_could_not_reach(monkeypatch):
    def unreachable(*args, **kw):
        raise osm.SourceError("All OpenStreetMap (Overpass) servers failed: timeout")
    monkeypatch.setattr(pipeline.osm, "search", unreachable)
    with pytest.raises(PipelineError) as err:
        pipeline.run(SearchParams(source="osm"))
    assert str(err.value) == ("Couldn't reach the map data service (OpenStreetMap), so no "
                              "leads were found.")
    assert "timeout" in err.value.detail


def test_skipped_steps_and_slow_steps(monkeypatch):
    assert web.skipped_steps(SearchParams()) == [1]                 # no Google or Yelp key
    monkeypatch.setenv("GOOGLE_PLACES_API_KEY", "AIzaFAKEKEYFORTESTS000000000000000000")
    assert web.skipped_steps(SearchParams()) == []
    assert web.skipped_steps(SearchParams(source="google")) == [2]
    job = {}
    progress = web._Progress(job)
    progress(Step("OpenStreetMap: (server 1 of 4)", MAP, 0, 1))
    assert not progress.slow()
    progress.step_started -= web.SLOW_SECONDS[2] + 1
    assert progress.slow()


# ---- marks and calls that can't be read are an error, never "not checked"

def test_unreadable_marks_are_an_error_not_unchecked(monkeypatch):
    lead = _lead()
    saved.save_search([lead])
    marks.set_mark(lead.uid, "yes")
    client = web.create_app().test_client()

    def broken(uids):
        raise RuntimeError("connection lost")
    monkeypatch.setattr(marks, "get_all", broken)
    for path in ("/leads", "/stats"):
        res = client.get(path)
        assert res.status_code == 503 and res.get_json()["error"] == web.MARKS_DOWN
    page = client.get("/download/saved.xlsx", headers={"Accept": "text/html"})
    assert page.status_code == 503 and b"Yes/No marks" in page.data
    assert b"Wright AI Solutions" in page.data and b"Back to Lead Finder" in page.data


def test_unreadable_calls_are_an_error_too(monkeypatch):
    saved.save_search([_lead()])

    def broken(leads):
        raise RuntimeError("connection lost")
    monkeypatch.setattr(calls, "apply", broken)
    res = web.create_app().test_client().get("/leads")
    assert res.status_code == 503 and "calls" in res.get_json()["error"]


def test_database_errors_never_show_class_names(monkeypatch):
    def down():
        raise store.Unavailable("the database could not be reached")
    monkeypatch.setattr(store, "open_db", down)
    client = web.create_app().test_client()
    for res in (client.get("/leads"), client.get("/stats"), client.get("/searches"),
                client.post("/mark", json={"by": "Sam", "key": "k", "value": "yes"}),
                client.post("/search", data={})):
        error = res.get_json()["error"]
        assert res.status_code == 503 and "Error" not in error and "(" not in error


# ---- the administrator's off switches

def test_search_pause_switch(monkeypatch):
    ran = []
    monkeypatch.setattr(web.finding, "run", lambda params, progress: ran.append(params) or RunResult(
        [], (40.76, -111.89), "SLC", [], {}))
    client = web.create_app().test_client()
    saved.save_search([_lead()])
    monkeypatch.setenv(config.SEARCH_PAUSED_ENV, "1")
    res = client.post("/search", data={})
    assert res.status_code == 503 and res.get_json()["error"] == web.SEARCH_PAUSED
    assert client.get("/searches").get_json()["paused"] == web.SEARCH_PAUSED
    assert b"Searching is paused" in client.get("/").data
    assert client.get("/leads").status_code == 200 and client.get("/stats").status_code == 200
    assert not ran and not client.get("/searches").get_json()["used_today"]
    monkeypatch.setenv(config.SEARCH_PAUSED_ENV, "")            # takes effect at once
    job = client.post("/search", data={}).get_json()["job_id"]
    assert _wait(client, job)["state"] == "done" and len(ran) == 1


def test_google_and_yelp_off_switches(monkeypatch):
    monkeypatch.setenv("GOOGLE_PLACES_API_KEY", "AIzaFAKEKEYFORTESTS000000000000000000")
    monkeypatch.setenv("YELP_API_KEY", "y" * 120)
    google, yelp_key, _ = SearchParams().resolved_keys()
    assert google and yelp_key
    monkeypatch.setenv(config.GOOGLE_OFF_ENV, "true")
    monkeypatch.setenv(config.YELP_OFF_ENV, "yes")
    assert SearchParams().resolved_keys()[:2] == ("", "")
    assert web.yelp_quota() is None
    for source in ("google", "both", "yelp"):
        with pytest.raises(PipelineError, match="switched off by the administrator"):
            pipeline.run(SearchParams(source=source))
    called = []
    monkeypatch.setattr(pipeline.google_places, "search", lambda *a, **k: called.append("g"))
    monkeypatch.setattr(pipeline.yelp, "search", lambda *a, **k: called.append("y"))
    monkeypatch.setattr(pipeline.osm, "search", lambda *a, **k: ([_lead()], []))
    res = pipeline.run(SearchParams())                 # map data still works
    assert not called and [lead.name for lead in res.leads] == ["Costco"]
    page = web.create_app().test_client().get("/").data
    assert b'value="google"' not in page and b'value="yelp"' not in page


# ---- error pages and logged-out visits

def test_branded_error_pages_and_json_for_the_api():
    client = web.create_app().test_client()
    page = client.get("/does-not-exist", headers={"Accept": "text/html"})
    assert page.status_code == 404 and b"Page not found" in page.data
    assert b"Wright AI Solutions" in page.data and b'href="/"' in page.data
    api = client.get("/does-not-exist")
    assert api.status_code == 404 and api.get_json()["error"]
    assert client.get("/status/nope").get_json()["error"]


def test_logged_out_browsers_go_to_the_login_page():
    client = web.create_app(password="s3cret", username="Matt").test_client()
    for path in ("/stats", "/does-not-exist"):
        res = client.get(path, headers={"Accept": "text/html,*/*;q=0.8"})
        assert res.status_code == 302 and res.headers["Location"].endswith("/login")
    assert client.get("/stats").status_code == 401          # the page's own data requests
    missing = client.get("/does-not-exist", headers={"Accept": "text/html"})
    assert missing.status_code == 302


def test_login_page_says_what_it_is_and_whom_to_ask(monkeypatch):
    client = web.create_app(password="s3cret", username="Matt").test_client()
    # Not set: the page still says whom to ask.
    page = client.get("/login").data.decode()
    assert "baler or compactor" in page
    assert "Ask the person who gave you your login, or Wright AI Solutions." in page
    # Set on Render: the named contact, the phone number and email ready to tap.
    monkeypatch.setenv("LEADGEN_SUPPORT_CONTACT", "Jane Doe at (801) 555-0100 or jane@example.com.")
    page = client.get("/login").data.decode()
    assert ('Contact Jane Doe at <a href="tel:8015550100">(801) 555-0100</a> or '
            '<a href="mailto:jane@example.com">jane@example.com</a>.</p>') in page
    assert "Wright AI Solutions." not in page.split("<footer>")[0]
    # Whatever is typed into the setting is shown as text, never as markup.
    monkeypatch.setenv("LEADGEN_SUPPORT_CONTACT", "<b>IT</b> desk, +1 801-555-0199")
    page = client.get("/login").data.decode()
    assert 'Contact &lt;b&gt;IT&lt;/b&gt; desk, <a href="tel:+18015550199">+1 801-555-0199</a>.' in page
    monkeypatch.setattr(web.time, "sleep", lambda s: None)
    wrong = client.post("/login", data={"username": "Matt", "password": "no"}).data.decode()
    assert 'type="password" autocomplete="current-password" required autofocus' in wrong


# ---- one database connection per page load

def test_leads_page_uses_one_database_connection(monkeypatch):
    saved.save_search([_lead(), _lead("Walmart", "w")])
    client = web.create_app().test_client()
    opened = []
    real = store.open_db
    monkeypatch.setattr(store, "open_db", lambda: opened.append(1) or real())
    for path in ("/leads", "/stats", "/download/saved.csv"):
        opened.clear()
        assert client.get(path).status_code == 200
        assert len(opened) == 1, path


def test_many_saved_leads_load_quickly():
    leads = [_lead(f"Business {i}", f"n{i}", lat=40.3 + (i // 60) * 0.01,
                   lon=-112.2 + (i % 60) * 0.01) for i in range(3000)]
    saved.save_search(leads, config.DEFAULT_KEYWORDS)
    client = web.create_app().test_client()
    start = time.time()
    assert len(client.get("/leads?limit=3000").get_json()["leads"]) == 3000
    assert time.time() - start < 5                  # ~0.3 s locally; generous for CI


# ---- scores don't drift with what later searches typed

def test_saved_scores_ignore_later_search_keywords():
    first = _lead("Acme Recycling Center", "r", raw_categories=["amenity=recycling"])
    saved.save_search([first])
    before = saved.load()[0]
    again = Lead(name="Acme Recycling Center", lat=40.72, lon=-111.9, source="osm",
                 source_id="r", raw_categories=["amenity=recycling"])
    score_lead(again, ["acme", "recycling center"])     # someone typed their own keywords
    assert again.score != before.score
    saved.save_search([again])
    after = saved.load()[0]
    assert (after.score, after.tier) == (before.score, before.tier)
    assert stats.summarize([after])["by_tier"] == stats.summarize([before])["by_tier"]


# ---- the Excel file

def test_excel_generated_time_is_utah_time_and_offline_columns_are_labelled(monkeypatch):
    import datetime as dt
    fixed = dt.datetime(2026, 9, 29, 19, 22, tzinfo=dt.UTC).timestamp()
    monkeypatch.setattr("leadgen.export.time.time", lambda: fixed)
    book = load_workbook(io.BytesIO(to_xlsx_bytes([_lead()], {})))
    info = {row[0].value: row[1].value for row in book["Run Info"].iter_rows() if row[0].value}
    assert info["Generated"] == "Sep 29, 2026, 1:22 PM (Utah time)"
    assert "Nothing typed there is saved in the website" in info["This file only"]
    header = [c.value for c in book["Leads"][1]]
    assert header[-2:] == ["My notes: equipment seen (this file only)", "My notes (this file only)"]
    assert "Has Baler or Compactor?" in header


def test_excel_of_10000_leads_is_quick_and_complete():
    """The Excel file is built in time that grows in step with the list: 10,000 saved
    leads in a few seconds (it took minutes once), with every column, the tier and
    competitor colours, the web links, the notes dropdown and Run Info."""
    from leadgen.export import COMPETITOR_FILL, OFFLINE_VERIFIED, TIER_FILLS, columns_for
    leads = []
    for i in range(10_000):
        lead = _lead(f"Store {i}", f"s{i}", website=f"https://store{i}.example.com",
                     map_url=f"https://www.openstreetmap.org/node/{i}", phone="8015550100",
                     city="Salt Lake City", address=f"{i} Main St", zip="84101")
        lead.distance_miles = 2.5
        leads.append(lead)
    leads[1].lead_type = "Competitor"
    started = time.perf_counter()
    data = to_xlsx_bytes(leads, {"List": "All saved leads"})
    assert time.perf_counter() - started < 3
    with zipfile.ZipFile(io.BytesIO(data)) as z:        # every row and link is there
        sheet = z.read("xl/worksheets/sheet1.xml").decode()
        assert sheet.count("<row ") == 10_001 and sheet.count("<hyperlink ") == 20_000
    # The rest is checked on a shorter list (reading 10,000 rows back takes a while).
    book = load_workbook(io.BytesIO(to_xlsx_bytes(leads[:200], {"List": "All saved leads"})))
    ws = book["Leads"]
    names = [name for name, _, _ in columns_for(leads[:200])]
    assert ws.max_row == 201 and ws.max_column == len(names)
    assert [c.value for c in ws[1]] == names
    assert ws["A2"].fill.fgColor.rgb == "FF" + TIER_FILLS[leads[0].tier]
    assert ws["D3"].fill.fgColor.rgb == "FF" + COMPETITOR_FILL          # a competitor row
    site = ws.cell(row=201, column=names.index("Website") + 1)
    where = ws.cell(row=201, column=names.index("Map Link") + 1)
    assert site.hyperlink.target == "https://store199.example.com"
    assert where.value == "Open map" and where.hyperlink.target.endswith("/node/199")
    last = ws.cell(row=1, column=len(names)).column_letter
    assert ws.freeze_panes == "F2" and ws.auto_filter.ref == f"A1:{last}{ws.max_row}"
    rule = ws.data_validations.dataValidation[0]
    letter = ws.cell(row=2, column=names.index(OFFLINE_VERIFIED) + 1).column_letter
    assert str(rule.sqref) == f"{letter}2:{letter}201"
    info = {row[0]: row[1] for row in book["Run Info"].iter_rows(values_only=True) if row[0]}
    assert info["List"] == "All saved leads" and info["Tiers"].startswith("A >= 60")


def test_saved_list_run_info_says_what_the_file_holds():
    from leadgen import marks, saved
    leads = [_lead("Costco", "c"), _lead("Harmons", "h", lat=40.8), _lead("Pro Baler", "p", lat=40.6)]
    saved.save_search(leads, config.DEFAULT_KEYWORDS)
    marks.set_mark(leads[0].uid, "yes")
    client = web.create_app().test_client()
    book = load_workbook(io.BytesIO(client.get("/download/saved.xlsx").data))
    info = {row[0]: row[1] for row in book["Run Info"].iter_rows(values_only=True) if row[0]}
    assert info["List"] == "All saved leads" and info["Leads in this file"] == 3
    assert info["Marked Yes (has a baler or compactor)"] == 1 and info["Marked No"] == 0
    assert info["Not checked yet"] == 1 and info["Competitors and Arco's own listing"] == 1
    assert sum(v for k, v in info.items() if k.startswith("Tier ") and isinstance(v, int)) == 3
    # No search on record (saved straight away): the times are when the leads were saved.
    assert info["First saved"].endswith("(Utah time)") and info["Latest saved"]
    assert "First search" not in info


def test_run_info_search_times_match_the_search_history(monkeypatch):
    day, _ = daily.claim({"location": "84104"})
    daily.finish(day, {"leads": 1})
    started = daily.history()["searches"][0]
    # The leads are saved minutes after the search started (a search takes a while).
    real = time.time
    with monkeypatch.context() as later:
        later.setattr(saved.time, "time", lambda: real() + 420)
        saved.save_search([_lead("Costco", "c")], config.DEFAULT_KEYWORDS)
    book = load_workbook(io.BytesIO(web.create_app().test_client().get("/download/saved.xlsx").data))
    info = {row[0]: row[1] for row in book["Run Info"].iter_rows(values_only=True) if row[0]}
    assert info["First search"] == f"{started['when']} (Utah time)" == info["Latest search"]


def test_one_place_for_utah_time():
    from leadgen import localtime
    assert localtime.clock_text(1790709540) == "1:19 PM"   # 19:19 UTC, daylight time
    assert daily.today() == localtime.now().strftime("%Y-%m-%d")


# ---- connections are kept between requests (store.py's pool)

@pytest.fixture
def new_connections(monkeypatch):
    """How many new database connections were opened, and how many times the tables
    were created, from now on."""
    made = {"connections": 0, "tables": 0}
    real_new, real_create = store._new_connection, store._create_tables

    def counted_new(key, postgres):
        made["connections"] += 1
        return real_new(key, postgres)

    def counted_create(db):
        made["tables"] += 1
        real_create(db)
    monkeypatch.setattr(store, "_new_connection", counted_new)
    monkeypatch.setattr(store, "_create_tables", counted_create)
    store.close_pool()
    return made


def test_ten_page_loads_reuse_one_connection(new_connections):
    saved.save_search([_lead(), _lead("Walmart", "w")])
    client = web.create_app().test_client()
    store.close_pool()
    new_connections["connections"] = 0
    for _ in range(10):
        assert client.get("/leads").status_code == 200
    assert new_connections["connections"] == 1


def test_the_tables_are_created_once_per_process(new_connections, monkeypatch):
    monkeypatch.setattr(store, "_ready", set())             # as if the server just started
    client = web.create_app().test_client()
    for _ in range(5):
        assert client.get("/leads").status_code == 200
        store.close_pool()                                  # every load on a new connection
    assert new_connections["connections"] >= 5 and new_connections["tables"] == 1


def test_the_pool_keeps_at_most_one_connection_per_server_thread(new_connections):
    import threading
    busy = [store.open_db() for _ in range(store.POOL_SIZE + 3)]   # a busy moment
    threads = [threading.Thread(target=store.release, args=(db,)) for db in busy]
    for t in threads:
        t.start()
    for t in threads:
        t.join()
    assert len(store._pool[busy[0].key]) == store.POOL_SIZE
    # The pool matches the web server's threads (render.yaml).
    import pathlib
    render = (pathlib.Path(__file__).resolve().parent.parent / "render.yaml").read_text()
    assert f"--threads {store.POOL_SIZE} " in render


def test_a_dropped_connection_is_replaced_without_an_error(new_connections, monkeypatch):
    saved.save_search([_lead()])
    client = web.create_app().test_client()
    assert client.get("/leads").status_code == 200
    (db, _), = store._pool[next(iter(store._pool))]
    db.conn.close()                                         # the database dropped it while idle
    assert client.get("/leads").status_code == 200
    # One idle a while is asked first; one that doesn't answer is replaced too.
    monkeypatch.setattr(store, "POOL_CHECK_AFTER", -1.0)
    asked = []
    real_one = store.Db.one

    def failing_ping(self, sql, params=()):
        if sql == "SELECT 1":
            asked.append(sql)
            raise OSError("server closed the connection")
        return real_one(self, sql, params)
    monkeypatch.setattr(store.Db, "one", failing_ping)
    before = new_connections["connections"]
    assert client.get("/leads").status_code == 200
    assert asked and new_connections["connections"] == before + 1
