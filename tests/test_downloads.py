"""The downloads (Excel and CSV) in plain words."""

import csv
import io

from leadgen import config
from leadgen.export import to_csv_bytes, to_xlsx_bytes
from leadgen.models import Lead
from leadgen.scoring import score_lead


def _map_lead(name, tags, sid=None, **kw):
    return Lead(name=name, lat=40.72, lon=-111.9, source="osm", source_id=sid or name,
                raw_categories=tags, **kw)


def test_downloads_after_a_free_search_use_plain_words():
    leads = [_map_lead("Acme Foods", ["industrial=food"], sources=["osm"]),
             _map_lead("Harmons", ["shop=supermarket"], sources=["osm"])]
    for lead in leads:
        score_lead(lead, config.DEFAULT_KEYWORDS)
    rows = list(csv.reader(io.StringIO(to_csv_bytes(leads).decode("utf-8-sig"))))
    head, body = rows[0], rows[1:]
    columns = {name: [row[i] for row in body] for i, name in enumerate(head)}
    assert columns["Sources"] == ["OpenStreetMap map data"] * 2
    assert "osm" not in " ".join(columns["Sources"])
    assert all("(from the map listing)" in why and "(by " not in why for why in columns["Why This Score"])
    for name in ("Found By", "Google Reviews", "Yelp Reviews", "Called By", "Call Notes"):
        assert name not in head
    # The columns kept are ones with something in them, or the core ones and the notes columns.
    from openpyxl import load_workbook
    ws = load_workbook(io.BytesIO(to_xlsx_bytes(leads)))["Leads"]
    assert [c.value for c in ws[1]] == head
    info = {r[0]: r[1] for r in load_workbook(io.BytesIO(to_xlsx_bytes(leads)))["Run Info"]
            .iter_rows(values_only=True) if r[0]}
    assert "Found By" in info["Columns left out"]


def test_download_file_names_carry_the_utah_date(monkeypatch):
    """Each download is named with the Utah date, so kept copies can be told apart."""
    from leadgen import daily, saved, web

    lead = _map_lead("Harmons", ["shop=supermarket"])
    score_lead(lead, config.DEFAULT_KEYWORDS)
    saved.save_search([lead])
    monkeypatch.setattr(daily, "today", lambda: "2026-09-30")
    client = web.create_app().test_client()
    for fmt in ("xlsx", "csv"):
        res = client.get(f"/download/saved.{fmt}")
        assert res.status_code == 200
        assert res.headers["Content-Disposition"] == \
            f"attachment; filename=compactor-leads-2026-09-30.{fmt}"


def test_no_rule_names_or_map_terms_in_the_downloads_or_on_the_leads_page():
    """A lead decided by each of the classifier's rules, saved: the Leads page's reasons
    and every cell of the Excel and CSV downloads are in a salesperson's words (no
    "category rule (...)", "map tag", "catch-all", rule texts...)."""
    import copy

    from test_classify_rules import RULE_EXAMPLES, internal_words

    from leadgen import saved, web
    leads = []
    for i, (_, example, _) in enumerate(RULE_EXAMPLES):
        lead = copy.deepcopy(example)
        lead.lat += i * 0.02                           # far apart: never merged as one site
        leads.append(score_lead(lead, config.DEFAULT_KEYWORDS))
    saved.save_search(leads, config.DEFAULT_KEYWORDS)
    client = web.create_app().test_client()
    shown = client.get("/leads?tab=all&limit=100").get_json()["leads"]
    assert len(shown) == len(RULE_EXAMPLES)
    for lead in shown:
        assert lead["reasons"] and not internal_words(" | ".join(lead["reasons"])), lead["reasons"]
    text = client.get("/download/saved.csv").data.decode("utf-8-sig")
    assert "Only the building type or the business name suggests what it does" in text
    from openpyxl import load_workbook
    book = load_workbook(io.BytesIO(client.get("/download/saved.xlsx").data))
    cells = [str(c) for ws in book.worksheets for row in ws.iter_rows(values_only=True) for c in row if c]
    for found in (internal_words(text), internal_words(" | ".join(cells))):
        assert not found, found


def test_last_called_is_a_real_date_in_excel_and_text_in_the_csv(monkeypatch):
    """Excel can sort and filter Last Called by date: each cell is a date and time on Utah's
    clock, shown as the site writes times; the CSV keeps the readable text."""
    import datetime as dt

    from openpyxl import load_workbook

    from leadgen import calls, saved, web
    from leadgen.localtime import date_time_text

    names = ("Harmons", "Smith's", "Costco", "Walmart")
    leads = [_map_lead(name, ["shop=supermarket"]) for name in names]
    for lead in leads:
        score_lead(lead, config.DEFAULT_KEYWORDS)
    saved.save_search(leads)
    # Called on these days (UTC): in summer (MDT) and in winter (MST), and after 6 PM Utah time,
    # which is already the next day in UTC.
    when = {"Harmons": dt.datetime(2026, 7, 3, 1, 5, tzinfo=dt.UTC),       # Jul 2, 7:05 PM MDT
            "Smith's": dt.datetime(2026, 1, 15, 18, 30, tzinfo=dt.UTC),    # Jan 15, 11:30 AM MST
            "Costco": dt.datetime(2026, 10, 2, 18, 39, tzinfo=dt.UTC)}     # Oct 2, 12:39 PM MDT
    for lead in leads:
        if lead.name in when:
            monkeypatch.setattr(calls.time, "time", lambda t=when[lead.name].timestamp(): t)
            calls.log_call(lead.uid, "Follow Up", "", by="Dana")
    client = web.create_app().test_client()
    ws = load_workbook(io.BytesIO(client.get("/download/saved.xlsx").data))["Leads"]
    head = [c.value for c in ws[1]]
    col = head.index("Last Called")
    cells = {row[head.index("Business Name")].value: row[col] for row in ws.iter_rows(min_row=2)}
    assert cells["Harmons"].value == dt.datetime(2026, 7, 2, 19, 5)
    assert cells["Smith's"].value == dt.datetime(2026, 1, 15, 11, 30)
    assert cells["Costco"].value == dt.datetime(2026, 10, 2, 12, 39)
    assert cells["Walmart"].value is None
    assert all(cells[n].is_date and cells[n].number_format == "mmm d, yyyy, h:mm AM/PM" for n in when)
    # Oldest to newest, as Excel's sort would put them.
    assert sorted(when, key=lambda n: cells[n].value) == ["Smith's", "Harmons", "Costco"]
    rows = {r["Business Name"]: r for r in csv.DictReader(io.StringIO(
        client.get("/download/saved.csv").data.decode("utf-8-sig")))}
    assert rows["Costco"]["Last Called"] == "Oct 2, 2026, 12:39 PM" == date_time_text(when["Costco"].timestamp())
    assert rows["Harmons"]["Last Called"] == "Jul 2, 2026, 7:05 PM" and rows["Walmart"]["Last Called"] == ""
