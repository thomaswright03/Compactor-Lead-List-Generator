import io
import json
from pathlib import Path

import pytest
from openpyxl import load_workbook

from leadgen import pipeline
from leadgen.export import format_phone, to_csv_bytes, to_xlsx_bytes
from leadgen.pipeline import PipelineError, SearchParams
from leadgen.sources import google_places, osm

FIX = Path(__file__).parent / "fixtures"


@pytest.fixture(autouse=True)
def no_cache(tmp_path, monkeypatch):
    monkeypatch.setattr("leadgen.http.CACHE_DIR", tmp_path / "cache")
    monkeypatch.delenv("GOOGLE_PLACES_API_KEY", raising=False)


def test_parse_google_place():
    data = json.loads((FIX / "google_page.json").read_text())
    lead = google_places.parse_place(data["places"][0], "supermarket")
    assert lead.name == "Smith's Marketplace"
    assert (lead.address, lead.city, lead.state, lead.zip) == ("455 S 500 E", "Salt Lake City", "UT", "84102")
    assert lead.rating_count == 2450 and lead.search_terms == ["supermarket"]
    bare = google_places.parse_place(data["places"][1])
    assert (bare.address, bare.city, bare.state, bare.zip) == ("1 Industrial Way", "Salt Lake City", "UT", "84104")


def _fake_google(monkeypatch, pages_per_query=3, fail=None):
    """Each query returns `pages_per_query` pages; fail(body, n) may raise HttpError."""
    calls = []

    def fake(method, url, **kw):
        body = kw["json_body"]
        calls.append(body)
        if fail:
            fail(body, len(calls))
        page = int(body.get("pageToken", "t0")[1:]) + 1
        place = {"id": f"{body['textQuery']}-{page}", "displayName": {"text": body["textQuery"]},
                 "location": {"latitude": 40.76, "longitude": -111.89}, "types": ["store"]}
        out = {"places": [place]}
        if page < pages_per_query:
            out["nextPageToken"] = f"t{page}"
        return out

    monkeypatch.setattr(google_places, "request_json", fake)
    monkeypatch.setattr(google_places, "TOKEN_DELAY_SECONDS", 0)
    return calls


def test_google_budget_is_spent_breadth_first(monkeypatch):
    calls = _fake_google(monkeypatch)
    leads, n, warnings = google_places.search(40.76, -111.89, 30, ["a", "b"], "KEY",
                                              grid_cells=1, max_requests=3)
    assert n == 3
    assert [(c["textQuery"], c.get("pageToken")) for c in calls] == [("a", None), ("b", None), ("a", "t1")]
    assert any("cap" in w for w in warnings)
    assert calls[0]["locationBias"]["circle"]["radius"] == 48280.32


def test_google_warns_when_cap_cannot_cover_every_search(monkeypatch):
    _fake_google(monkeypatch, pages_per_query=1)
    _, n, warnings = google_places.search(40.76, -111.89, 30, list("abcde"), "KEY",
                                          grid_cells=7, max_requests=10)
    assert n == 10
    assert any("35 searches needed" in w for w in warnings)


def test_google_page_token_error_keeps_results(monkeypatch):
    from leadgen.http import HttpError

    def fail(body, n):
        if body.get("pageToken"):
            raise HttpError("returned HTTP 400: INVALID_ARGUMENT")
    _fake_google(monkeypatch, fail=fail)
    leads, _, warnings = google_places.search(40.76, -111.89, 30, ["a", "b"], "KEY")
    assert len(leads) == 2
    assert any("later results page failed" in w for w in warnings)


def test_google_bad_key_raises(monkeypatch):
    from leadgen.http import HttpError

    def fail(body, n):
        raise HttpError("returned HTTP 400: API key not valid")
    _fake_google(monkeypatch, fail=fail)
    with pytest.raises(google_places.SourceError):
        google_places.search(40.76, -111.89, 30, ["a"], "KEY")


def test_google_caches_complete_chains_only(monkeypatch):
    calls = _fake_google(monkeypatch)
    google_places.search(40.76, -111.89, 30, ["a", "b"], "KEY", max_requests=4)
    assert len(calls) == 4          # a: 2 of 3 pages, b: 2 of 3 pages -> neither complete
    calls.clear()
    leads, n, _ = google_places.search(40.76, -111.89, 30, ["a", "b"], "KEY")
    assert n == 6 and len(leads) == 6   # re-fetched from page 1, no stale tokens
    calls.clear()
    leads, n, _ = google_places.search(40.76, -111.89, 30, ["a", "b"], "KEY")
    assert n == 0 and len(calls) == 0 and len(leads) == 6   # served from cache


def test_pipeline_searches_keywords_and_competitors_first(monkeypatch):
    seen = {}

    def fake_search(lat, lon, radius, queries, *a, **k):
        seen["queries"] = queries
        return [], 0, []
    monkeypatch.setattr(google_places, "search", fake_search)
    monkeypatch.setattr(osm, "search", lambda *a, **k: ([], []))
    pipeline.run(SearchParams(source="both", api_key="KEY", keywords=["baler", "compactor"]))
    assert seen["queries"][:4] == ["baler", "compactor", "Pro Baler", "Action Compaction"]


def test_osm_query_and_parse():
    q = osm.build_query(40.76, -111.89, 30, ["baler"])
    assert "around:48280" in q and "baler" in q and "out tags bb" in q
    data = json.loads((FIX / "overpass.json").read_text())
    leads = [l for l in map(osm.parse_element, data["elements"]) if l]
    assert len(leads) == 4  # unnamed node dropped
    wh = next(l for l in leads if l.name == "Acme Distribution Center")
    assert wh.footprint_sqft > 100_000 and "building=warehouse" in wh.raw_categories
    assert "brand=Smith's" in leads[0].raw_categories


def _fake_sources(monkeypatch):
    gdata = json.loads((FIX / "google_page.json").read_text())
    odata = json.loads((FIX / "overpass.json").read_text())
    monkeypatch.setattr(google_places, "search", lambda *a, **k: (
        [google_places.parse_place(p, "q") for p in gdata["places"]], 1, []))
    monkeypatch.setattr(osm, "search", lambda *a, **k: (
        [l for l in map(osm.parse_element, odata["elements"]) if l], []))


def test_pipeline_end_to_end(monkeypatch):
    _fake_sources(monkeypatch)
    res = pipeline.run(SearchParams(source="both", api_key="KEY", min_score=20))
    names = [l.name for l in res.leads]
    assert "Far Away Foods" not in names            # outside 30 miles
    assert "Old Closed Warehouse" not in names      # permanently closed
    assert names.count("Smith's Marketplace") == 1  # google + osm merged
    assert names[0] == "Smith's Marketplace"        # highest score first
    comps = {l.name for l in res.leads if l.lead_type == "Competitor"}
    assert comps == {"Pro Baler", "Action Compaction Services"}  # kept despite low score
    assert res.stats["competitors flagged"] == 2


def test_pipeline_osm_only_without_key(monkeypatch):
    _fake_sources(monkeypatch)
    res = pipeline.run(SearchParams())
    assert all("osm" in l.sources for l in res.leads)
    assert any("No Google Places or Yelp API key" in w for w in res.warnings)


def test_pipeline_google_without_key_errors():
    with pytest.raises(PipelineError):
        pipeline.run(SearchParams(source="google"))


def test_exports(monkeypatch):
    _fake_sources(monkeypatch)
    params = SearchParams(source="both", api_key="KEY")
    res = pipeline.run(params)
    csv_text = to_csv_bytes(res.leads).decode("utf-8-sig")
    assert csv_text.splitlines()[0].startswith("Score,Tier,Lead Type,Flags,Business Name")
    assert "(801) 328-1683" in csv_text
    wb = load_workbook(io.BytesIO(to_xlsx_bytes(res.leads, res.run_info(params))))
    assert wb.sheetnames == ["Leads", "Run Info"]
    assert wb["Leads"].max_row == len(res.leads) + 1
    info = {r[0]: r[1] for r in wb["Run Info"].iter_rows(values_only=True) if r and r[0]}
    assert "api_key" not in info


def test_format_phone():
    assert format_phone("+1 801-555-0100") == "(801) 555-0100"
    assert format_phone("ext 12") == "ext 12"


def test_osm_skips_roads_labels_and_campus_footprints():
    road = {"type": "way", "id": 1, "bounds": {"minlat": 40, "minlon": -112, "maxlat": 40.01, "maxlon": -111.99},
            "tags": {"name": "University Pkwy N", "highway": "primary"}}
    label = {"type": "way", "id": 2, "bounds": {"minlat": 40, "minlon": -112, "maxlat": 40.001, "maxlon": -111.999},
             "tags": {"name": "B", "building": "apartments"}}
    resort = {"type": "relation", "id": 3,
              "bounds": {"minlat": 40.6, "minlon": -111.6, "maxlat": 40.65, "maxlon": -111.5},
              "tags": {"name": "Big Resort", "building": "hotel", "tourism": "hotel"}}
    assert osm.parse_element(road) is None
    assert osm.parse_element(label) is None
    assert osm.parse_element(resort).footprint_sqft is None
    q = osm.build_query(40.76, -111.89, 30)
    assert '[!"highway"]' in q and '[!"type"]' not in q


def test_pipeline_limit_keeps_competitors(monkeypatch):
    _fake_sources(monkeypatch)
    res = pipeline.run(SearchParams(source="both", api_key="KEY", limit=1))
    types = sorted(l.lead_type for l in res.leads)
    assert types == ["Competitor", "Competitor", "Prospect"]


def test_pipeline_drops_place_google_says_closed_even_with_osm_copy(monkeypatch):
    from leadgen.models import Lead
    g = Lead(name="Old Mill Foods", lat=40.75, lon=-111.9, source="google", source_id="g",
             raw_categories=["point_of_interest"], business_status="CLOSED_PERMANENTLY")
    o = Lead(name="Old Mill Foods", lat=40.7501, lon=-111.9001, source="osm", source_id="n",
             raw_categories=["industrial=food"])
    monkeypatch.setattr(google_places, "search", lambda *a, **k: ([g], 1, []))
    monkeypatch.setattr(osm, "search", lambda *a, **k: ([o], []))
    assert pipeline.run(SearchParams(source="both", api_key="KEY")).leads == []
    kept = pipeline.run(SearchParams(source="both", api_key="KEY", include_closed=True)).leads
    assert len(kept) == 1


def test_pipeline_rejects_bad_grid():
    with pytest.raises(PipelineError):
        pipeline.run(SearchParams(grid=2))


def test_osm_escaping_and_skips():
    q = osm.build_query(40.76, -111.89, 30, ["c++", 'say "hi"', "a.b"])
    assert r"c\\+\\+" in q and r'say\\ \"hi\"' in q and r"a\\.b" in q
    assert "restaurant" not in q and "clothes" not in q   # classify-only tags are not fetched
    hist = {"type": "node", "id": 9, "lat": 40.7, "lon": -111.9,
            "tags": {"name": "Saint Marks Hospital (historical)", "amenity": "hospital"}}
    gone = {"type": "node", "id": 10, "lat": 40.7, "lon": -111.9,
            "tags": {"name": "Old Market", "disused:shop": "supermarket"}}
    assert osm.parse_element(hist) is None and osm.parse_element(gone) is None


def test_exports_neutralize_formulas_and_control_chars():
    from leadgen.models import Lead
    evil = Lead(name='=HYPERLINK("http://x","click")', lat=40.7, lon=-111.9, source="osm",
                source_id="n", address="@SUM(A1)", city="\x01=1+41", website="javascript:alert(1)",
                primary_category="Bad\x07Tag", tier="C", score=20)
    csv_text = to_csv_bytes([evil]).decode("utf-8-sig")
    assert "'=HYPERLINK" in csv_text and "'@SUM" in csv_text and "'=1+41" in csv_text
    wb = load_workbook(io.BytesIO(to_xlsx_bytes([evil])))
    ws = wb["Leads"]
    values = [c.value for c in ws[2]]
    assert values[4] == '=HYPERLINK("http://x","click")'            # shown as text...
    assert all(c.data_type != "f" for row in ws.iter_rows() for c in row)   # ...never run
    assert "Bad\x07Tag" not in values and "BadTag" in values
    assert ws.cell(row=2, column=12).hyperlink is None   # javascript: is never linked


def test_http_errors_redact_keys():
    from leadgen.http import redact
    assert "SECRET" not in redact("GET /geocode/json?address=x&key=SECRET failed")


def test_google_grid_cap_still_searches_every_phrase(monkeypatch):
    calls = _fake_google(monkeypatch, pages_per_query=1)
    google_places.search(40.76, -111.89, 30, list("abcde"), "KEY", grid_cells=7, max_requests=5)
    assert {c["textQuery"] for c in calls} == set("abcde")


def test_google_transient_first_failure_then_bad_key(monkeypatch):
    from leadgen.http import HttpError

    def fail(body, n):
        if n == 1:
            raise HttpError("returned HTTP 503")
        raise HttpError("returned HTTP 400: API key not valid. API_KEY_INVALID")
    _fake_google(monkeypatch, fail=fail)
    with pytest.raises(google_places.SourceError):
        google_places.search(40.76, -111.89, 30, ["a", "b", "c"], "KEY")


def test_closed_google_pin_just_outside_radius_closes_inside_copy(monkeypatch):
    from leadgen.geo import offset_point
    from leadgen.models import Lead
    olat, olon = offset_point(40.7608, -111.8910, 30.02, 90)
    ilat, ilon = offset_point(40.7608, -111.8910, 29.99, 90)
    g = Lead(name="Old Mill Foods", lat=olat, lon=olon, source="google", source_id="g",
             business_status="CLOSED_PERMANENTLY")
    o = Lead(name="Old Mill Foods", lat=ilat, lon=ilon, source="osm", source_id="n",
             raw_categories=["industrial=food"])
    monkeypatch.setattr(google_places, "search", lambda *a, **k: ([g], 1, []))
    monkeypatch.setattr(osm, "search", lambda *a, **k: ([o], []))
    assert pipeline.run(SearchParams(source="both", api_key="KEY")).leads == []


def test_osm_keeps_live_places_with_history_and_3m():
    was = {"type": "node", "id": 1, "lat": 40.7, "lon": -111.9,
           "tags": {"name": "Acme Foods", "industrial": "food", "was:name": "Old Acme"}}
    three_m = {"type": "node", "id": 2, "lat": 40.7, "lon": -111.9,
               "tags": {"name": "3M", "man_made": "works"}}
    depot = {"type": "way", "id": 3, "bounds": {"minlat": 40.7, "minlon": -111.9, "maxlat": 40.701, "maxlon": -111.899},
             "tags": {"name": "UTA Depot Warehouse", "building": "warehouse", "railway": "yard"}}
    assert osm.parse_element(was) and osm.parse_element(three_m) and osm.parse_element(depot)


@pytest.mark.parametrize("extra", [{}, {"limit": 1}, {"include_closed": True},
                                   {"keywords": ["baler"], "only_keyword_matches": True}])
def test_search_numbers_add_up(monkeypatch, extra):
    """The search history's Details read as a funnel whose numbers add up."""
    from leadgen.web.finding import plain_details
    _fake_sources(monkeypatch)
    res = pipeline.run(SearchParams(source="both", api_key="KEY", min_score=20, **extra))
    s = res.stats
    assert s["results in radius"] - s["duplicates merged"] == s["after dedupe"]
    assert (s["after dedupe"] - s["closed"] - s["below min score"] - s["not matching keywords"]
            - s["over limit"]) == s["leads kept"]
    assert sum(s[f"tier {t}"] for t in "ABCD") == s["leads kept"]
    assert s["leads kept"] == sum(l.business_status != "CLOSED_PERMANENTLY" for l in res.leads)
    assert s["closed"] == 1 and s["min score"] == 20
    labels = [label for label, _ in plain_details(s)]
    assert labels.index("Listings within the radius") < labels.index("Businesses after merging duplicates") \
        < labels.index("Left out: score below 20") < labels.index("Leads kept") < labels.index("Tier D leads")
    assert labels[0].startswith("Businesses from") and labels[-1] == "Took"
