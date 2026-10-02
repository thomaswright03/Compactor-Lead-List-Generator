"""A lead without a street address or city still shows a town (and ZIP): the one its map
position is near, worked out offline and labelled "near ...", in the pages and the
downloads, so same-named stores can be told apart. Nothing saved is changed."""

import csv
import io
import json

from openpyxl import load_workbook

from leadgen import config, places, saved, store, web
from leadgen.models import Lead
from leadgen.scoring import score_lead


def _lowes(sid, lat, lon, **kw):
    lead = Lead(name="Lowe's", lat=lat, lon=lon, source="osm", source_id=sid,
                raw_categories=["shop=doityourself"], **kw)
    return score_lead(lead, config.DEFAULT_KEYWORDS)


def test_the_nearest_town_and_zip_come_from_the_coordinates():
    assert places.near(*config.OWN_COORDS) == places.Near("Salt Lake City", "UT", "84104")
    assert places.near(40.6097, -111.9391).text() == "near West Jordan, UT 84088"
    assert places.near(40.5649, -111.8389).town == "Sandy"
    assert places.near(41.06, -111.97).town == "Layton"
    assert places.near(None, None) is None and places.near(0.0, 0.0) is None
    assert places.near(25.76, -80.19) is None            # Miami: not near any town in the table


def _three_lowes():
    leads = [_lowes("w1", 40.6097, -111.9391), _lowes("s1", 40.5649, -111.8389),
             _lowes("m1", 40.6669, -111.8880, address="5000 S State St", city="Murray", zip="84107")]
    saved.save_search(leads, config.DEFAULT_KEYWORDS)
    return leads


def test_the_leads_page_shows_which_town_each_store_is_near():
    _three_lowes()
    client = web.create_app().test_client()
    rows = client.get("/leads?tab=all&q=lowe").get_json()["leads"]
    near = sorted(r["near"] for r in rows)
    assert near == ["", "near Sandy, UT 84094", "near West Jordan, UT 84088"]
    assert all(r["city"] in ("", "Murray") for r in rows)        # never passed off as the city
    # The filter and the city sort know the town too.
    assert [r["near"] for r in client.get("/leads?tab=all&q=west%20jordan").get_json()["leads"]] == \
        ["near West Jordan, UT 84088"]
    by_city = client.get("/leads?tab=all&sort=city&dir=asc").get_json()["leads"]
    assert [r["city"] or r["near"] for r in by_city] == ["Murray", "near Sandy, UT 84094",
                                                        "near West Jordan, UT 84088"]
    # The saved rows themselves are untouched.
    with store.connect() as db:
        stored = [json.loads(lead) for (lead,) in db.all("SELECT lead FROM leads")]
    assert all("near" not in lead for lead in stored)
    assert sorted(lead["city"] for lead in stored) == ["", "", "Murray"]


def test_every_downloaded_row_with_coordinates_has_a_town():
    _three_lowes()
    client = web.create_app().test_client()
    rows = list(csv.DictReader(io.StringIO(client.get("/download/saved.csv").data.decode("utf-8-sig"))))
    assert [r for r in rows if r["Latitude"] and not r["City"]] == []
    cities = sorted((r["City"], r["State"], r["ZIP"]) for r in rows)
    assert cities == [("Murray", "", "84107"), ("near Sandy", "UT", "near 84094"),
                      ("near West Jordan", "UT", "near 84088")]
    book = load_workbook(io.BytesIO(client.get("/download/saved.xlsx").data))
    info = {row[0]: row[1] for row in book["Run Info"].iter_rows(values_only=True) if row[0]}
    assert "worked out from the map position" in info["City / ZIP “near …”"]
