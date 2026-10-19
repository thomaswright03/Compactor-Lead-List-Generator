"""Sign-up mode (LEADGEN_SIGNUP=1): the public landing page and its plans, the Terms and
Privacy written for many companies, Yelp left out, the Overpass server setting and the
map data's credit in the downloads. With the flag off the site is AARCO's, unchanged."""

import csv
import io
import re
from pathlib import Path

import pytest
from openpyxl import load_workbook

from leadgen import config, http, pipeline, plans, signup_mode, web
from leadgen.export import OSM_CREDIT, to_csv_bytes, to_xlsx_bytes
from leadgen.models import Lead
from leadgen.pipeline import PipelineError, SearchParams
from leadgen.sources import osm
from tests.test_html_valid import html_problems

PAGE = {"Accept": "text/html"}
YELP_KEY = "y" * 128
GOOGLE_KEY = "AIza" + "x" * 35


@pytest.fixture
def signup(monkeypatch):
    monkeypatch.setenv("LEADGEN_SIGNUP", "1")


def _text(page):
    """The page without its styles and scripts (the theme's comments name AARCO's colours)."""
    return re.sub(r"<(style|script)\b.*?</\1>", "", page, flags=re.S)


def _site():
    return web.create_app(password="secret", username="Matt").test_client()


@pytest.mark.parametrize(("value", "on"), [("", False), ("0", False), ("no", False),
                                           ("1", True), ("true", True), (" YES ", True), ("on", True)])
def test_the_flag_is_read_from_the_environment(monkeypatch, value, on):
    monkeypatch.setenv("LEADGEN_SIGNUP", value)
    assert signup_mode.on() is on


def test_the_plans_numbers_are_in_one_place():
    starter, team = plans.PLANS["starter"], plans.PLANS["team"]
    assert (starter.name, starter.monthly_price, starter.users, starter.areas,
            starter.searches_per_month, starter.lookups_per_month) == ("Starter", 99, 1, 1, 8, 500)
    assert (team.name, team.monthly_price, team.users, team.areas,
            team.searches_per_month, team.lookups_per_month) == ("Team", 249, 5, 3, 30, 2500)
    assert plans.TRIAL_DAYS == 14 and plans.CARD_REQUIRED and plans.AREA_MAX_MILES == 30
    assert (plans.annual_price(starter), plans.annual_price(team)) == (990, 2490)   # 2 months free


def test_without_the_flag_there_is_no_landing_page():
    assert web.create_app().test_client().get("/welcome", headers=PAGE).status_code == 404
    # Behind the login it is like any other address: off to the login page.
    res = _site().get("/welcome", headers=PAGE)
    assert res.status_code == 302 and res.headers["Location"].endswith("/login")


def test_the_landing_page_opens_without_logging_in_and_links_on(signup):
    res = _site().get("/welcome", headers=PAGE)
    assert res.status_code == 200
    page = res.get_data(as_text=True)
    assert html_problems(page) == []
    assert "Lead Finder" in page and "balers, compactors and waste equipment" in page
    assert ("See every business in your area likely to need what you service, scored and on a map. "
            "Mark who has it, log every call, and hide the customers you already have.") in page
    for link in ('href="/signup"', 'href="/login"', 'href="/terms"', 'href="/privacy"'):
        assert link in page
    assert ">Start free trial</a>" in page and ">Log in</a>" in page
    assert "How it works" in page and "never leaves your computer" in page
    for figure in ("$99", "$249", "8 searches a month", "30 searches a month", "500 phone lookups",
                   "2,500 phone lookups", "1 user<", "5 users", "3 areas", "up to 30 miles",
                   "14-day free trial", "card required", "2 months free", "Custom"):
        assert figure in page, figure
    assert "AARCO" not in _text(page)
    # Phone-friendly like the other pages: the viewport, 16 px gutters, 44 px targets.
    assert '<meta name="viewport" content="width=device-width, initial-scale=1">' in page
    assert "padding:0 var(--sp-4)" in page and "min-height:44px" in page
    assert "prefers-color-scheme: dark" in page                     # the theme (_theme.html)


def test_the_landing_page_follows_the_plans(signup, monkeypatch):
    monkeypatch.setitem(plans.PLANS, "starter", plans.Plan("starter", "Starter", 109, 1, 1, 9, 600))
    page = web.create_app().test_client().get("/welcome").get_data(as_text=True)
    assert "$109" in page and "9 searches a month" in page and "$1,090 a year" in page


def test_the_legal_pages_differ_by_mode(monkeypatch):
    client = _site()
    aarco = {p: client.get(p).get_data(as_text=True) for p in ("/terms", "/privacy")}
    assert "AARCO" in aarco["/terms"] and "Terms of use" in aarco["/terms"]
    assert "CAN-SPAM" not in aarco["/terms"] and "salted" not in aarco["/privacy"]
    monkeypatch.setenv("LEADGEN_SIGNUP", "1")
    terms = client.get("/terms")
    privacy = client.get("/privacy")
    assert terms.status_code == privacy.status_code == 200                  # still without a login
    terms, privacy = terms.get_data(as_text=True), privacy.get_data(as_text=True)
    for page in (terms, privacy):
        assert "AARCO" not in _text(page) and html_problems(page) == []
        assert "https://www.openstreetmap.org/copyright" in page
    for words in ("Terms of service", "Subscription, trial and cancellation", "14-day free trial",
                  "card is required", "cancel at any time", "owns the data", "export",
                  "on the customer's behalf", "do-not-call", "CAN-SPAM", "Google Maps", "State of Utah",
                  "Current customer import"):
        assert words in terms, words
    for words in ("never uploaded", "salted with a value unique to your company", "test a guess",
                  "Clear the customer list", "one-way hashes", "kept separate", "encrypted (AES-256)",
                  "export all of its data or delete it"):
        assert words in privacy, words
    monkeypatch.setenv("LEADGEN_SIGNUP", "")
    assert {p: client.get(p).get_data(as_text=True) for p in aarco} == aarco


def test_yelp_is_off_in_signup_mode_even_with_a_key(signup, monkeypatch):
    monkeypatch.setenv("YELP_API_KEY", YELP_KEY)
    monkeypatch.setenv("GOOGLE_PLACES_API_KEY", GOOGLE_KEY)
    google, yelp_key, _ = SearchParams().resolved_keys()
    assert (google, yelp_key) == (GOOGLE_KEY, "")
    # A Yelp key in the Google setting isn't used for Yelp either.
    monkeypatch.delenv("YELP_API_KEY")
    assert SearchParams(api_key=YELP_KEY).resolved_keys() == ("", "", True)
    # Find leads: Google and the map data, never Yelp; asking for Yelp alone is refused.
    assert pipeline._sources_to_use(SearchParams(), GOOGLE_KEY, "", False) == (True, False, True)
    assert pipeline._sources_to_use(SearchParams(), "", "", False) == (False, False, True)
    with pytest.raises(PipelineError, match="needs YELP_API_KEY"):
        pipeline.run(SearchParams(source="yelp"))
    monkeypatch.setenv("YELP_API_KEY", YELP_KEY)
    page = web.create_app().test_client().get("/").get_data(as_text=True)
    assert '<option value="yelp">' not in page and '<option value="google">' in page
    assert web.yelp_quota() is None
    # Without the flag the same key is used, as on AARCO's site.
    monkeypatch.setenv("LEADGEN_SIGNUP", "0")
    assert SearchParams().resolved_keys()[1] == YELP_KEY


def test_the_overpass_server_can_be_set(monkeypatch):
    assert config.overpass_endpoints() == config.OVERPASS_ENDPOINTS        # default: the public mirrors
    tried = []

    def answer(method, url, **kw):
        tried.append(url)
        return {"elements": []}
    monkeypatch.setattr(osm, "request_json", answer)
    own = "https://overpass.example.com/api/interpreter"
    monkeypatch.setenv("LEADGEN_OVERPASS_URL", f" {own} , https://backup.example.com/api/interpreter")
    assert config.overpass_endpoints() == [own, "https://backup.example.com/api/interpreter"]
    assert osm.search(40.76, -111.89, 5) == ([], [])
    assert tried == [own]
    assert http.source_of(own) == "osm"                       # counted as the map data (health.py)


def _leads():
    return [Lead(name="Harmons", lat=40.72, lon=-111.9, source="osm", source_id="h",
                 raw_categories=["shop=supermarket"], sources=["osm"])]


def test_downloads_credit_openstreetmap_only_in_signup_mode(monkeypatch):
    plain_csv, plain_xlsx = to_csv_bytes(_leads()), to_xlsx_bytes(_leads())
    assert b"OpenStreetMap contributors" not in plain_csv
    info = {r[0]: r[1] for r in load_workbook(io.BytesIO(plain_xlsx))["Run Info"].iter_rows(values_only=True)
            if r[0]}
    assert "Data sources" not in info
    monkeypatch.setenv("LEADGEN_SIGNUP", "1")
    rows = list(csv.reader(io.StringIO(to_csv_bytes(_leads()).decode("utf-8-sig"))))
    assert rows[-2:] == [[], [OSM_CREDIT]] and "© OpenStreetMap contributors" in OSM_CREDIT
    assert rows[:-2] == list(csv.reader(io.StringIO(plain_csv.decode("utf-8-sig"))))
    info = {r[0]: r[1] for r in load_workbook(io.BytesIO(to_xlsx_bytes(_leads())))["Run Info"]
            .iter_rows(values_only=True) if r[0]}
    assert info["Data sources"] == OSM_CREDIT
    assert "https://www.openstreetmap.org/copyright" in OSM_CREDIT


def test_the_map_credits_openstreetmap():
    """The Map page's street map carries the credit itself (Leaflet's attribution)."""
    script = (Path(web.STATIC) / "map.js").read_text()
    assert '&copy; <a href="https://www.openstreetmap.org/copyright">OpenStreetMap</a> ' in script
    assert '"contributors"' in script and "attribution: TILE_CREDIT" in script
