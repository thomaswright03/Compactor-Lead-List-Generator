"""Round 8: search words choose what is searched (scores stay on the standard words),
campus units and neighbours are not campus prospects, the Yelp reset names its day,
plain progress and downloads, one note for an incomplete search, and the in-site
emergency switches."""

import csv
import io
import json
import time
from pathlib import Path

import pytest

from leadgen import config, dedupe, pipeline, saved, switches, web
from leadgen.export import to_csv_bytes, to_xlsx_bytes
from leadgen.localtime import day_clock_text
from leadgen.models import Lead
from leadgen.pipeline import SearchParams
from leadgen.scoring import score_lead
from leadgen.sources import osm


def _wait(client, job):
    for _ in range(400):
        body = client.get(f"/status/{job}").get_json()
        if body["state"] != "running":
            return body
        time.sleep(0.05)
    raise AssertionError("the search never finished")


def _map_lead(name, tags, sid=None, **kw):
    return Lead(name=name, lat=40.72, lon=-111.9, source="osm", source_id=sid or name,
                raw_categories=tags, **kw)


# ---- extra search words: the saved and shown score is the one the minimum applies to

def _pallet_search(monkeypatch, **params):
    monkeypatch.setattr(pipeline, "geocode", lambda location, key: (40.72, -111.9, "Salt Lake City"))
    monkeypatch.setattr(osm, "search", lambda *a, **k: (
        [_map_lead("Acme Pallet Co", ["building=industrial"])], []))
    return pipeline.run(SearchParams(source="osm", **params))


def test_search_words_do_not_change_the_score(monkeypatch):
    plain = _pallet_search(monkeypatch).leads[0]
    worded = _pallet_search(monkeypatch, keywords=["pallet"]).leads[0]
    assert worded.score == plain.score == 28
    assert not any("pallet" in r for r in worded.reasons)
    assert worded.matched_keywords == ["pallet"]          # still says which words matched


def test_the_minimum_score_applies_to_the_score_that_is_saved(monkeypatch):
    left_out = _pallet_search(monkeypatch, keywords=["pallet"], min_score=30)
    assert left_out.leads == [] and left_out.stats["below min score"] == 1
    kept = _pallet_search(monkeypatch, keywords=["pallet"], min_score=20)
    saved.save_search(kept.leads, ["pallet"])
    shown = saved.load()
    assert [(l.name, l.score) for l in shown] == [("Acme Pallet Co", 28)]
    assert all(l.score >= 20 for l in shown)
    assert shown[0].reasons == kept.leads[0].reasons       # "Why this score" agrees too


def test_only_matching_the_search_words_still_filters(monkeypatch):
    kept = _pallet_search(monkeypatch, keywords=["pallet"], only_keyword_matches=True)
    assert [l.name for l in kept.leads] == ["Acme Pallet Co"]
    gone = _pallet_search(monkeypatch, keywords=["lumber"], only_keyword_matches=True)
    assert gone.leads == [] and gone.stats["not matching keywords"] == 1


def test_the_form_says_what_search_words_do():
    page = web.create_app().test_client().get("/").get_data(as_text=True)
    assert "score higher" not in page
    assert "scores always use the standard words" in page


# ---- a university's units and neighbours are not campus prospects

REFERENCE = json.loads((Path(__file__).parent / "fixtures" / "scoring_reference.json").read_text())


@pytest.mark.parametrize("item", REFERENCE["not_campus"], ids=lambda i: i["name"])
def test_campus_units_and_neighbours_are_not_campus_prospects(item):
    lead = _map_lead(item["name"], item["raw_categories"])
    score_lead(lead, config.DEFAULT_KEYWORDS)
    assert lead.category_key != "education" and lead.tier == "D"


@pytest.mark.parametrize("name,tags", [("University of Utah", ["amenity=university"]),
                                       ("Salt Lake Community College", []),
                                       ("College of Eastern Utah", []),
                                       ("Westminster College", ["university"])])
def test_campuses_are_still_campus_prospects(name, tags):
    lead = _map_lead(name, tags)
    score_lead(lead, config.DEFAULT_KEYWORDS)
    assert lead.category_key == "education" and lead.tier == "C"


# ---- the Yelp reset names its day

def test_a_time_names_today_or_tomorrow_in_utah():
    import datetime as dt
    noon = dt.datetime(2026, 9, 29, 18, 0, tzinfo=dt.UTC).timestamp()        # 12:00 PM in Utah
    assert day_clock_text(noon + 3600, noon) == "today at 1:00 PM"
    assert day_clock_text(noon + 24 * 3600, noon) == "tomorrow at 12:00 PM"
    assert day_clock_text(noon + 3 * 24 * 3600, noon) == "Oct 2 at 12:00 PM"
    # 11 PM in Utah is already the next day in UTC: still "today" in Utah.
    late = dt.datetime(2026, 9, 30, 5, 0, tzinfo=dt.UTC).timestamp()
    assert day_clock_text(late, noon) == "today at 11:00 PM"


# ---- progress and downloads in plain words

def test_progress_never_names_servers():
    job = {}
    progress = web._Progress(job)
    progress("OpenStreetMap: searching the free map data (server 1 of 4)")
    assert job["message"] == "Searching the free map data…"
    progress("OpenStreetMap: searching the free map data (server 3 of 4)")
    assert "server" not in job["message"] and "trying another source" in job["message"]


def test_downloads_after_a_free_search_use_plain_words():
    leads = [_map_lead("Acme Foods", ["industrial=food"], sources=["osm"]),
             _map_lead("Harmons", ["shop=supermarket"], sources=["osm"])]
    for lead in leads:
        score_lead(lead, config.DEFAULT_KEYWORDS)
    rows = list(csv.reader(io.StringIO(to_csv_bytes(leads).decode("utf-8-sig"))))
    head, body = rows[0], rows[1:]
    columns = {name: [row[i] for row in body] for i, name in enumerate(head)}
    assert columns["Sources"] == ["OpenStreetMap map data"] * 2
    assert "osm" not in " ".join(columns["Sources"])
    assert all("(from the map listing)" in why and "(by " not in why for why in columns["Why This Score"])
    for name in ("Found By", "Google Reviews", "Yelp Reviews", "Called By", "Call Notes"):
        assert name not in head
    # The columns kept are ones with something in them, or the core ones and the notes columns.
    from openpyxl import load_workbook
    ws = load_workbook(io.BytesIO(to_xlsx_bytes(leads)))["Leads"]
    assert [c.value for c in ws[1]] == head
    info = {r[0]: r[1] for r in load_workbook(io.BytesIO(to_xlsx_bytes(leads)))["Run Info"]
            .iter_rows(values_only=True) if r[0]}
    assert "Found By" in info["Columns left out"]


def test_merging_is_a_public_function_of_dedupe():
    assert callable(dedupe.merge)
    assert "_merge" not in Path(saved.__file__).read_text()


# ---- an incomplete search shows one note, the details behind "More"

def test_an_incomplete_search_gives_one_short_note(monkeypatch):
    monkeypatch.setattr(pipeline, "geocode", lambda location, key: (40.72, -111.9, "Salt Lake City"))

    def partial(*a, **k):
        lead = _map_lead("Costco", ["shop=wholesale"])
        raise osm.PartialResult("OpenStreetMap answered for only part of the area", [lead], [])
    monkeypatch.setattr(pipeline.osm, "search", partial)
    client = web.create_app().test_client()
    job = client.post("/search", data={"location": "84101"}).get_json()["job_id"]
    body = _wait(client, job)
    assert body["note"] == ("Some businesses are missing: the map data service answered for "
                            "only part of the area. Run the search again today to fill them in.")
    assert len(body["warnings"]) >= 2                     # the details, shown behind "More"


# ---- the emergency switches work from inside the site, with no restart

def test_pausing_in_the_site_stops_the_next_search():
    client = web.create_app().test_client()
    assert client.post("/switches", json={"key": "search_paused", "on": True, "by": "Tom"}).status_code == 200
    refused = client.post("/search", data={"location": "84101"})
    assert refused.status_code == 503 and "Searching is paused" in refused.get_json()["error"]
    listed = client.get("/searches").get_json()
    assert listed["paused"] and listed["switches"]["search_paused"]["by"] == "Tom"
    assert listed["switches"]["search_paused"]["when"]
    assert config.stop_reason() == "the administrator paused searching"   # a running search stops
    client.post("/switches", json={"key": "search_paused", "on": False})
    assert client.get("/searches").get_json()["paused"] is None
    assert config.stop_reason() is None


def test_paid_sources_can_be_switched_off_in_the_site():
    client = web.create_app().test_client()
    client.post("/switches", json={"key": "yelp_off", "on": True})
    assert config.stop_reason("yelp") == "the administrator switched Yelp off"
    assert config.stop_reason("google") is None
    with pytest.raises(pipeline.PipelineError, match="Yelp searches are switched off"):
        pipeline.run(SearchParams(source="yelp", yelp_api_key="k"))


def test_the_environment_switch_stays_as_a_backup(monkeypatch):
    monkeypatch.setenv(config.SEARCH_PAUSED_ENV, "1")
    client = web.create_app().test_client()
    client.post("/switches", json={"key": "search_paused", "on": False})
    state = client.get("/searches").get_json()["switches"]["search_paused"]
    assert state["on"] and state["env"]
    assert switches.is_on(config.SEARCH_PAUSED_ENV)


def test_an_unknown_switch_is_refused():
    client = web.create_app().test_client()
    assert client.post("/switches", json={"key": "everything", "on": True}).status_code == 400
    assert client.post("/switches", json={"key": "yelp_off", "on": "yes"}).status_code == 400
