"""Verified phone numbers and contact names: saved on a lead by a salesperson, with who
saved them, shown beside the listing's own phone on the page and in the downloads, and
never changed by a later search."""

import csv
import io

import pytest
from openpyxl import load_workbook

from leadgen import config, contacts, saved, store, web
from leadgen.models import Lead
from leadgen.scoring import score_lead


def _lead(name="Galleria Mall", sid="way/227957948", **kw):
    kw.setdefault("lat", 40.72)
    kw.setdefault("lon", -111.9)
    kw.setdefault("raw_categories", ["shop=mall"])
    lead = Lead(name=name, source="osm", source_id=sid, **kw)
    score_lead(lead, config.DEFAULT_KEYWORDS)
    return lead


def _save(client, uid, phone="", contact="", by="Dana"):
    return client.post("/contact", json={"key": uid, "phone": phone, "contact": contact, "by": by})


def test_a_verified_phone_and_contact_are_saved_with_who_saved_them():
    lead = _lead()
    saved.save_search([lead])
    client = web.create_app().test_client()
    answer = _save(client, lead.uid, "801.555.0199", "  Jane Doe,\n facilities manager ")
    assert answer.status_code == 200
    contact = answer.get_json()["contact"]
    assert contact["verified_phone"] == "(801) 555-0199"
    assert contact["contact_name"] == "Jane Doe, facilities manager"
    assert contact["contact_by"] == "Dana" and contact["contact_when"]
    [row] = client.get("/leads").get_json()["leads"]
    assert row["phone"] == "" and row["verified_phone"] == "(801) 555-0199"
    assert row["contact_name"] == "Jane Doe, facilities manager" and row["contact_by"] == "Dana"
    history = client.get(f"/calls/{lead.uid}").get_json()["contacts"]
    assert [(c["phone"], c["contact"], c["by"]) for c in history] == [
        ("(801) 555-0199", "Jane Doe, facilities manager", "Dana")]


@pytest.mark.parametrize(("phone", "ok"), [
    ("(801) 555-0199", True), ("+1 801 555 0199", True), ("801-555-0199 ext. 12", True),
    ("8015550199", True), ("555-0199", False), ("call the front desk", False), ("1234567890123", False)])
def test_a_verified_phone_needs_its_area_code(phone, ok):
    lead = _lead()
    saved.save_search([lead])
    answer = _save(web.create_app().test_client(), lead.uid, phone)
    assert (answer.status_code == 200) is ok
    if not ok:
        assert answer.get_json()["error"] == contacts.BAD_PHONE


def test_saving_needs_a_name_a_saved_lead_and_something_typed():
    lead = _lead()
    saved.save_search([lead])
    client = web.create_app().test_client()
    assert _save(client, lead.uid, "801-555-0199", by=" ").status_code == 400
    assert _save(client, "nope", "801-555-0199").status_code == 400
    nothing = _save(client, lead.uid)
    assert nothing.status_code == 400 and nothing.get_json()["error"] == contacts.NOTHING
    other_site = client.post("/contact", json={"key": lead.uid, "phone": "801-555-0199", "by": "Dana"},
                             headers={"Origin": "https://example.com"})
    assert other_site.status_code == 403
    assert contacts.latest([lead.uid]) == {}


def test_every_save_is_kept_and_the_latest_shows():
    lead = _lead()
    saved.save_search([lead])
    client = web.create_app().test_client()
    _save(client, lead.uid, "801-555-0199", "Jane Doe", by="Dana")
    _save(client, lead.uid, "801-555-0199", "Jane Doe", by="Lee")       # the same again: no new row
    _save(client, lead.uid, "801-555-0123", "Jane Doe", by="Lee")
    [row] = client.get("/leads").get_json()["leads"]
    assert row["verified_phone"] == "(801) 555-0123" and row["contact_by"] == "Lee"
    # Both boxes emptied: taken off the lead, and the history keeps every save.
    cleared = _save(client, lead.uid, by="Sam").get_json()["contact"]
    assert cleared == {"verified_phone": "", "contact_name": "", "contact_by": "", "contact_when": "",
                       "contact_at": None, "contact_saves": 3}
    [row] = client.get("/leads").get_json()["leads"]
    assert row["verified_phone"] == "" and row["contact_name"] == "" and row["contact_by"] == ""
    assert row["contact_saves"] == 3          # the page offers its earlier versions
    history = client.get(f"/calls/{lead.uid}").get_json()["contacts"]
    assert [(c["phone"], c["by"]) for c in history] == [
        ("", "Sam"), ("(801) 555-0123", "Lee"), ("(801) 555-0199", "Dana")]
    with store.connect() as db:
        assert db.one("SELECT COUNT(*) FROM lead_contacts")[0] == 3


def test_a_later_search_never_changes_the_verified_contact():
    lead = _lead()
    saved.save_search([lead])
    client = web.create_app().test_client()
    _save(client, lead.uid, "801-555-0199", "Jane Doe")
    # Another search finds the business again, its listing now with a phone of its own.
    again = _lead(phone="+1 801-555-0100", address="580 Main Street")
    saved.save_search([again])
    assert again.uid == lead.uid
    [row] = client.get("/leads").get_json()["leads"]
    assert row["phone"] == "(801) 555-0100" and row["address"] == "580 Main Street"
    assert row["verified_phone"] == "(801) 555-0199" and row["contact_name"] == "Jane Doe"
    # And the stored lead holds only the listing's details.
    with store.connect() as db:
        stored = db.one("SELECT lead FROM leads WHERE uid = ?", (lead.uid,))[0]
    assert "555-0199" not in stored and "Jane Doe" not in stored
    file = list(csv.DictReader(io.StringIO(client.get("/download/saved.csv").data.decode("utf-8-sig"))))
    assert file[0]["Phone"] == "(801) 555-0100" and file[0]["Verified Phone"] == "(801) 555-0199"


def test_the_downloads_show_it_beside_the_listings_phone():
    a, b = _lead(), _lead("Costco Wholesale", "way/c", raw_categories=["shop=wholesale"], phone="801-555-0100")
    saved.save_search([a, b])
    client = web.create_app().test_client()
    _save(client, a.uid, "801-555-0199", "Jane Doe, facilities", by="Dana")
    rows = list(csv.reader(io.StringIO(client.get("/download/saved.csv").data.decode("utf-8-sig"))))
    head = rows[0]
    at = head.index("Phone")
    assert head[at:at + 4] == ["Phone", "Verified Phone", "Contact Name", "Verified By"]
    by_name = {r[head.index("Business Name")]: r for r in rows[1:]}
    assert by_name["Galleria Mall"][at:at + 3] == ["", "(801) 555-0199", "Jane Doe, facilities"]
    assert by_name["Galleria Mall"][at + 3].startswith("Dana, ")
    assert by_name["Costco Wholesale"][at:at + 4] == ["(801) 555-0100", "", "", ""]
    ws = load_workbook(io.BytesIO(client.get("/download/saved.xlsx").data))["Leads"]
    head = [c.value for c in ws[1]]
    at = head.index("Phone")
    assert head[at:at + 4] == ["Phone", "Verified Phone", "Contact Name", "Verified By"]
    cells = {r[head.index("Business Name")].value: [c.value for c in r[at:at + 3]] for r in ws.iter_rows(min_row=2)}
    assert cells["Galleria Mall"] == [None, "(801) 555-0199", "Jane Doe, facilities"]


def test_the_columns_are_left_out_when_nobody_verified_anything():
    lead = _lead()
    saved.save_search([lead])
    head = next(csv.reader(io.StringIO(web.create_app().test_client().get(
        "/download/saved.csv").data.decode("utf-8-sig"))))
    assert not {"Verified Phone", "Contact Name", "Verified By"} & set(head)


def test_a_verified_phone_counts_for_has_phone_and_the_filter_finds_the_contact():
    a, b = _lead(), _lead("Costco Wholesale", "way/c", raw_categories=["shop=wholesale"])
    saved.save_search([a, b])
    client = web.create_app().test_client()
    assert client.get("/leads?phone=1").get_json()["leads"] == []
    _save(client, a.uid, "801-555-0199", "Jane Doe")
    assert [l["key"] for l in client.get("/leads?phone=1").get_json()["leads"]] == [a.uid]
    assert [l["key"] for l in client.get("/leads?q=jane").get_json()["leads"]] == [a.uid]


def test_open_pages_hear_of_a_colleagues_save():
    lead = _lead()
    saved.save_search([lead])
    client = web.create_app().test_client()
    now = client.get("/leads").get_json()["now"]
    _save(client, lead.uid, "801-555-0199")
    changed = client.get(f"/leads?since={now}").get_json()["leads"]
    assert [l["key"] for l in changed] == [lead.uid] and changed[0]["verified_phone"] == "(801) 555-0199"


def test_merging_saved_rows_keeps_every_verified_contact(monkeypatch):
    first, second = _lead("Shoreline Ridge 825", "way/825", raw_categories=["building=apartments"]), \
        _lead("Shoreline Ridge 826", "way/826", raw_categories=["building=apartments"], lat=40.7203)
    with monkeypatch.context() as m:
        m.setattr(saved, "same_site", lambda a, b: False)
        m.setattr(saved, "is_duplicate", lambda a, b: False)
        saved.save_search([first])
        saved.save_search([second])
    assert len(saved.load()) == 2
    client = web.create_app().test_client()
    _save(client, second.uid, "801-555-0199", "Jane Doe")
    assert saved.merge_sites() == 1
    [row] = client.get("/leads").get_json()["leads"]
    assert row["key"] == first.uid and row["verified_phone"] == "(801) 555-0199"
    assert len(client.get(f"/calls/{first.uid}").get_json()["contacts"]) == 1
