"""The command line: `python -m leadgen run` writes Excel and CSV files, and Yelp
is only used when its calls count against the website's shared daily limit."""

import csv
import io
import os

import pytest
from openpyxl import load_workbook

from leadgen import cli, config, envfile, pipeline
from leadgen.models import Lead
from leadgen.sources import yelp


@pytest.fixture(autouse=True)
def in_tmp(tmp_path, monkeypatch):
    # No .env of the project is ever read by these tests.
    monkeypatch.chdir(tmp_path)


def _fake_osm(monkeypatch):
    def search(lat, lon, radius, keywords=(), progress=None):
        if progress:
            progress("OpenStreetMap: searching the free map data (server 1 of 4)")
        return [Lead(name="Smith's Marketplace", lat=lat + 0.01, lon=lon, source="osm",
                     source_id="node/1", raw_categories=["shop=supermarket"], city="Salt Lake City"),
                Lead(name="Pro Baler", lat=lat, lon=lon + 0.01, source="osm", source_id="node/2",
                     raw_categories=["craft=metal_construction"])], []
    monkeypatch.setattr(pipeline.osm, "search", search)


def test_run_writes_an_excel_file(tmp_path, monkeypatch, capsys):
    _fake_osm(monkeypatch)
    out = tmp_path / "out" / "leads.xlsx"
    assert cli.main(["run", "--location", "40.76,-111.89", "--radius", "10", "--out", str(out)]) == 0
    book = load_workbook(io.BytesIO(out.read_bytes()))
    names = [row[4] for row in book["Leads"].iter_rows(min_row=2, values_only=True)]
    assert "Smith's Marketplace" in names and "Pro Baler" in names     # competitors are kept
    printed = capsys.readouterr()
    assert "leads near" in printed.out and "Top leads:" in printed.out
    assert "searching the free map data" in printed.err                 # progress goes to stderr


def test_run_writes_a_csv_file_quietly(tmp_path, monkeypatch, capsys):
    _fake_osm(monkeypatch)
    out = tmp_path / "leads.csv"
    assert cli.main(["run", "-l", "40.76,-111.89", "-r", "10", "-q", "--min-score", "0",
                     "--limit", "1", "-k", "compactor,baler", "--out", str(out)]) == 0
    rows = list(csv.reader(io.StringIO(out.read_bytes().decode("utf-8-sig"))))
    assert rows[0][:5] == ["Score", "Tier", "Lead Type", "Flags", "Business Name"]
    assert len(rows) == 3                     # the top prospect, plus the competitor
    assert "searching the free map data" not in capsys.readouterr().err    # no progress


def test_an_other_file_type_becomes_excel(tmp_path, monkeypatch):
    _fake_osm(monkeypatch)
    assert cli.main(["run", "-l", "40.76,-111.89", "-q", "--out", str(tmp_path / "leads.txt")]) == 0
    assert (tmp_path / "leads.xlsx").read_bytes()[:2] == b"PK"


def test_a_failed_search_exits_2_with_the_reason(monkeypatch, capsys):
    def down(*args, **kw):
        raise pipeline.SourceError("All OpenStreetMap (Overpass) servers failed: timeout")
    monkeypatch.setattr(pipeline.osm, "search", down)
    assert cli.main(["run", "-l", "40.76,-111.89", "-q"]) == 2
    err = capsys.readouterr().err
    assert "Error: Couldn't reach the map data service" in err and "Details: All OpenStreetMap" in err


def test_no_command_prints_help(capsys):
    assert cli.main([]) == 1
    assert "usage: leadgen" in capsys.readouterr().out


def test_yelp_needs_the_sites_database(monkeypatch, capsys):
    """Without DATABASE_URL, Yelp calls would be counted in a local file, apart from
    the website's count of the shared key's 50 calls a day: the CLI refuses."""
    monkeypatch.setenv("YELP_API_KEY", "test-yelp-key-" + "x" * 40)
    monkeypatch.delenv("DATABASE_URL", raising=False)
    called = []
    monkeypatch.setattr(yelp, "search", lambda *a, **k: called.append(1) or ([], 0, []))
    assert cli.main(["run", "-l", "40.76,-111.89", "--source", "yelp", "-q"]) == 2
    err = capsys.readouterr().err
    assert "DATABASE_URL" in err and "50" in err and not called


def test_auto_leaves_yelp_out_without_the_sites_database(monkeypatch, capsys):
    monkeypatch.setenv("YELP_API_KEY", "test-yelp-key-" + "x" * 40)
    monkeypatch.delenv("DATABASE_URL", raising=False)
    _fake_osm(monkeypatch)
    called = []
    monkeypatch.setattr(yelp, "search", lambda *a, **k: called.append(1) or ([], 0, []))
    assert cli.main(["run", "-l", "40.76,-111.89", "-q", "--out", "leads.csv"]) == 0
    assert not called
    assert "Yelp was left out" in capsys.readouterr().err


def test_yelp_is_used_with_the_sites_database(monkeypatch):
    """With DATABASE_URL (the website's database), the CLI shares the site's count."""
    monkeypatch.setenv("YELP_API_KEY", "test-yelp-key-" + "x" * 40)
    monkeypatch.setenv("DATABASE_URL", os.environ.get("LEADGEN_TEST_DATABASE_URL") or "postgresql://unused")
    _fake_osm(monkeypatch)
    called = []
    monkeypatch.setattr(yelp, "search", lambda *a, **k: called.append(1) or ([], 0, []))
    assert cli.main(["run", "-l", "40.76,-111.89", "-q", "--out", "leads.csv"]) == 0
    assert called


def test_dotenv_fills_only_missing_settings(tmp_path, monkeypatch):
    (tmp_path / ".env").write_text("# a comment\nLEADGEN_T_ONE='first'\nLEADGEN_T_TWO=second\n"
                                   "not a setting\nLEADGEN_T_EMPTY=\n")
    monkeypatch.setenv("LEADGEN_T_TWO", "kept")
    monkeypatch.delenv("LEADGEN_T_ONE", raising=False)
    envfile.load_dotenv()
    assert os.environ["LEADGEN_T_ONE"] == "first" and os.environ["LEADGEN_T_TWO"] == "kept"
    assert "LEADGEN_T_EMPTY" not in os.environ
    monkeypatch.delenv("LEADGEN_T_ONE")
    envfile.load_dotenv(str(tmp_path / "missing.env"))          # no file: nothing happens


def test_the_yelp_limit_is_the_sites(capsys):
    with pytest.raises(SystemExit):
        cli.main(["run", "--help"])
    assert f"{config.YELP_DAILY_LIMIT} calls" in capsys.readouterr().out


def test_reference_copies_the_sites_marks(tmp_path, capsys):
    import json

    from leadgen import marks, saved
    from leadgen.scoring import score_lead
    leads = [Lead(name=n, lat=40.7 + i / 100, lon=-111.9, source="yelp", source_id=f"y{i}", phone="8015550100",
                  raw_categories=cats, yelp_reviews=200)
             for i, (n, cats) in enumerate([("Harmons", ["yelp:grocery"]), ("Nail Spa", ["yelp:othersalons"]),
                                            ("Pro Baler", ["yelp:junkremoval"]), ("Hampton Inn", ["yelp:hotels"])])]
    for lead in leads:
        score_lead(lead, config.DEFAULT_KEYWORDS)
    saved.save_search(leads, config.DEFAULT_KEYWORDS)
    marks.set_mark(leads[0].uid, "yes")
    marks.set_mark(leads[1].uid, "no")
    marks.set_mark(leads[2].uid, "yes")                  # a competitor: not a prospect
    out = tmp_path / "ref.json"
    out.write_text(json.dumps({"about": "x", "businesses": []}))
    assert cli.main(["reference", "--out", str(out)]) == 0
    data = json.loads(out.read_text())
    assert data["about"] == "x" and data["businesses"] == []
    assert [(e["name"], e["marked"]) for e in data["confirmed"]] == [("Nail Spa", "no"), ("Harmons", "yes")]
    assert "phone" not in data["confirmed"][0] and data["confirmed"][1]["tier"] == leads[0].tier
    assert "1 businesses marked Yes and 1 marked No" in capsys.readouterr().out
