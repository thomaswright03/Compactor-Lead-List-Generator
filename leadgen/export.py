"""CSV and Excel output."""

import csv
import io
import re
from datetime import datetime

from openpyxl import Workbook
from openpyxl.cell.cell import ILLEGAL_CHARACTERS_RE
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation

from .scoring import TIER_LABELS

COLUMNS = [
    ("Score", lambda l: l.score, 8),
    ("Tier", lambda l: TIER_LABELS.get(l.tier, l.tier), 13),
    ("Lead Type", lambda l: l.lead_type, 20),
    ("Flags", lambda l: "; ".join(l.flags), 28),
    ("Business Name", lambda l: l.name, 36),
    ("Category", lambda l: l.category, 30),
    ("Address", lambda l: l.address, 30),
    ("City", lambda l: l.city, 18),
    ("State", lambda l: l.state, 7),
    ("ZIP", lambda l: l.zip, 8),
    ("Phone", lambda l: format_phone(l.phone), 16),
    ("Website", lambda l: l.website, 32),
    ("Distance (mi)", lambda l: l.distance_miles, 12),
    ("Why This Score", lambda l: "Points: " + " | ".join(l.reasons) if l.reasons else "", 60),
    ("Matched Keywords", lambda l: ", ".join(l.matched_keywords), 18),
    ("Google Reviews", lambda l: l.rating_count, 10),
    ("Yelp Reviews", lambda l: l.yelp_reviews, 10),
    ("Approx. Footprint (sq ft)", lambda l: l.footprint_sqft, 14),
    ("Source Category", lambda l: l.primary_category, 22),
    ("Found By", lambda l: ", ".join(l.search_terms), 24),
    ("Sources", lambda l: ", ".join(l.sources or [l.source]), 12),
    ("Map Link", lambda l: l.map_url, 30),
    ("Latitude", lambda l: round(l.lat, 6), 11),
    ("Longitude", lambda l: round(l.lon, 6), 11),
    ("Has Baler?", lambda l: {"yes": "Yes", "no": "No"}.get(l.has_baler, ""), 11),
    ("Verified?", lambda l: "", 11),
    ("Notes", lambda l: "", 30),
]

TIER_FILLS = {"A": "C6EFCE", "B": "E2EFDA", "C": "FFF2CC", "D": "F2F2F2"}
COMPETITOR_FILL = "F8CBAD"


def format_phone(phone):
    digits = re.sub(r"\D", "", phone or "")
    if len(digits) == 11 and digits.startswith("1"):
        digits = digits[1:]
    if len(digits) == 10:
        return f"({digits[:3]}) {digits[3:6]}-{digits[6:]}"
    return (phone or "").strip()


_FORMULA_PREFIXES = ("=", "+", "-", "@", "\t", "\r")


def _safe(value):
    """CSV: neutralize text a spreadsheet would run as a formula (data comes from the public)."""
    if isinstance(value, str):
        value = ILLEGAL_CHARACTERS_RE.sub("", value)
    if isinstance(value, str) and value.startswith(_FORMULA_PREFIXES):
        return "'" + value
    return value


def _xl(value):
    """Safe for Excel: no control characters openpyxl rejects."""
    return ILLEGAL_CHARACTERS_RE.sub("", value) if isinstance(value, str) else value


def _as_text(row):
    """openpyxl stores strings starting with '=' as formulas; keep them as text."""
    for cell in row:
        if cell.data_type == "f":
            cell.data_type = "s"
            cell.quotePrefix = True   # stays text even if someone edits the cell


def rows(leads):
    return [[fn(l) for _, fn, _ in COLUMNS] for l in leads]


def to_csv_bytes(leads):
    buf = io.StringIO()
    w = csv.writer(buf)
    w.writerow([c for c, _, _ in COLUMNS])
    for r in rows(leads):
        w.writerow(["" if v is None else _safe(v) for v in r])
    return buf.getvalue().encode("utf-8-sig")   # BOM so Excel opens it cleanly


def to_xlsx_bytes(leads, run_info=None):
    wb = Workbook()
    ws = wb.active
    ws.title = "Leads"
    ws.append([c for c, _, _ in COLUMNS])
    for cell in ws[1]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor="1F4E78")
        cell.alignment = Alignment(vertical="center", wrap_text=True)
    for lead, r in zip(leads, rows(leads)):
        ws.append([_xl(v) for v in r])
        row = ws.max_row
        _as_text(ws[row])
        fill = COMPETITOR_FILL if lead.lead_type == "Competitor" else TIER_FILLS.get(lead.tier)
        if fill:
            for col in range(1, 5):
                ws.cell(row=row, column=col).fill = PatternFill("solid", fgColor=fill)
        site = ws.cell(row=row, column=_col("Website"))
        if re.match(r"https?://", lead.website or "", re.I):
            site.hyperlink = ILLEGAL_CHARACTERS_RE.sub("", lead.website)
            site.font = Font(color="0563C1", underline="single")
        link = ws.cell(row=row, column=_col("Map Link"))
        if re.match(r"https?://", lead.map_url or "", re.I):
            link.hyperlink = ILLEGAL_CHARACTERS_RE.sub("", lead.map_url)
            link.value = "Open map"
            link.font = Font(color="0563C1", underline="single")
    for i, (_, _, width) in enumerate(COLUMNS, start=1):
        ws.column_dimensions[get_column_letter(i)].width = width
    ws.freeze_panes = "F2"
    ws.auto_filter.ref = ws.dimensions
    if leads:
        dv = DataValidation(type="list", formula1='"Yes,No,Maybe,Has compactor,Has baler"',
                            allow_blank=True)
        col = get_column_letter(_col("Verified?"))
        dv.add(f"{col}2:{col}{ws.max_row}")
        ws.add_data_validation(dv)

    info = wb.create_sheet("Run Info")
    info.append(["Generated", datetime.now().strftime("%Y-%m-%d %H:%M")])
    for k, v in (run_info or {}).items():
        info.append([_xl(k), _xl(v if isinstance(v, (int, float, str)) else str(v))])
        _as_text(info[info.max_row])
    info.append([])
    info.append(["Tiers", "A >= 60, B >= 40, C >= 20, D below 20"])
    info.append(["Row colors", "Green = stronger lead, orange = competitor"])
    info.column_dimensions["A"].width = 22
    info.column_dimensions["B"].width = 90

    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def _col(name):
    return next(i for i, (c, _, _) in enumerate(COLUMNS, start=1) if c == name)
